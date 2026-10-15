"""Build data/roles/team_roles.xlsx: every Yahoo player on his current team with last season's process
stats, computed unit and two roles (offense + defense), plus Unit/Role 1/Role 2 2026-27 and Notes columns to fill
in by hand.

Usage:
  python scripts/build_roles.py            # pull 2025-26 tracking, starts and matchups (~38 nba_api calls), rebuild
  python scripts/build_roles.py --cached   # reuse data/stats/tracking/, just rebuild the sheet
Re-running keeps anything already typed into the input columns (matched by yahoo_player_id). A sheet in the old
depth/role/option layout is migrated once (see roles.migrate_legacy) and its untouched roles re-seeded.
Needs data/stats/season_lines/2025-26_regular.csv (scripts/pull_stats.py), data/adp.csv (scripts/pull_adp.py) and
data/stats/pbp/2025-26/ (scripts/pull_pbp.py) for the finisher signal.
"""
import argparse
import sys

import pandas as pd
from nba_api.stats.endpoints import (leaguedashplayershotlocations, leaguedashplayerstats, leaguedashptdefend,
                                     leaguedashptstats, leaguehustlestatsplayer, leagueseasonmatchups)
from nba_api.stats.static import teams
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from yahoo_fantasy.adp import load_adp
from yahoo_fantasy.id_map.store import read_map
from yahoo_fantasy.nba_client import nba_frame
from yahoo_fantasy.paths import DATA_DIR
from yahoo_fantasy.roles import (INJURED, INPUT_COLUMNS, ROLES, SHEET_COLUMNS, UNITS, build_table, carry_over, classify,
                                 derive_signals, is_legacy, migrate_legacy, poa_minutes, seed_roles, shape_tracking,
                                 shot_totals)
from yahoo_fantasy.stats.cache import STATS_DIR, dataset_path, write_dataset
from yahoo_fantasy.stats.transform import REGULAR

SEASON = "2025-26"
PBP_DIR = STATS_DIR / "pbp" / SEASON
MIN_PBP_GAMES = 1200
SHEET_PATH = DATA_DIR / "roles" / "team_roles.xlsx"
MIN_TRACKING_ROWS = 400
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
TRACKING_EXTRA = ["CATCH_SHOOT_FGA", "RIM_DFGA", "DEFLECTIONS", "POA_TOTAL"]


def pull_tracking() -> pd.DataFrame:
    def pt(measure: str):
        return nba_frame(leaguedashptstats.LeagueDashPtStats, timeout=60, season=SEASON,
                         season_type_all_star=REGULAR, player_or_team="Player", pt_measure_type=measure,
                         per_mode_simple="PerGame")
    shots = nba_frame(leaguedashplayershotlocations.LeagueDashPlayerShotLocations, timeout=60, season=SEASON,
                      season_type_all_star=REGULAR, per_mode_detailed="PerGame", distance_range="By Zone")
    starts = nba_frame(leaguedashplayerstats.LeagueDashPlayerStats, timeout=60, season=SEASON,
                       season_type_all_star=REGULAR, per_mode_detailed="Totals", starter_bench_nullable="Starters")
    gs = starts[["PLAYER_ID", "GP"]].rename(columns={"PLAYER_ID": "nba_player_id", "GP": "GS"})
    possessions = pt("Possessions")
    catch_shoot = pt("CatchShoot")[["PLAYER_ID", "CATCH_SHOOT_FGA"]].rename(columns={"PLAYER_ID": "nba_player_id"})
    rim = nba_frame(leaguedashptdefend.LeagueDashPtDefend, timeout=60, season=SEASON, season_type_all_star=REGULAR,
                    defense_category="Less Than 6Ft", per_mode_simple="PerGame")
    rim = rim[["CLOSE_DEF_PERSON_ID", "FGA_LT_06"]].rename(
        columns={"CLOSE_DEF_PERSON_ID": "nba_player_id", "FGA_LT_06": "RIM_DFGA"})
    hustle = nba_frame(leaguehustlestatsplayer.LeagueHustleStatsPlayer, timeout=60, season=SEASON,
                       season_type_all_star=REGULAR, per_mode_time="PerGame")
    hustle = hustle[["PLAYER_ID", "DEFLECTIONS"]].rename(columns={"PLAYER_ID": "nba_player_id"})
    ball_share = (possessions["TIME_OF_POSS"] / possessions["MIN"].where(possessions["MIN"] > 0)).set_axis(
        possessions["PLAYER_ID"])
    poa = poa_minutes(pull_matchups(), ball_share)
    out = shape_tracking(possessions, pt("Passing"), shots).merge(gs, on="nba_player_id", how="left")
    for extra in (catch_shoot, rim, hustle, poa):
        out = out.merge(extra, on="nba_player_id", how="left")
    return out


def pull_matchups() -> pd.DataFrame:
    """Season matchup totals, one call per defending team."""
    frames = [nba_frame(leagueseasonmatchups.LeagueSeasonMatchups, timeout=90, season=SEASON,
                        season_type_playoffs=REGULAR, per_mode_simple="Totals", def_team_id_nullable=str(t["id"]))
              for t in teams.get_teams()]
    return pd.concat(frames, ignore_index=True)[["OFF_PLAYER_ID", "DEF_PLAYER_ID", "MATCHUP_TIME_SEC"]]


def load_shot_totals() -> pd.DataFrame:
    files = sorted(PBP_DIR.glob("*.csv"))
    if len(files) < MIN_PBP_GAMES:
        sys.exit(f"ABORT: only {len(files)} play-by-play games in {PBP_DIR} (< {MIN_PBP_GAMES}); run scripts/pull_pbp.py")
    cols = ["personId", "isFieldGoal", "shotDistance", "shotResult", "subType", "description"]
    return shot_totals(pd.concat((pd.read_csv(f, usecols=cols) for f in files), ignore_index=True))


def read_previous(path) -> pd.DataFrame | None:
    if not path.exists():
        return None
    ws = load_workbook(path, read_only=True)["Roles"]
    rows = list(ws.values)
    headers = {label: col for col, label in SHEET_COLUMNS}
    return pd.DataFrame(rows[1:], columns=[headers.get(h, h) for h in rows[0]])


def write_sheet(table: pd.DataFrame, path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Roles"
    ws.append([label for _, label in SHEET_COLUMNS])
    for row in table.itertuples(index=False):
        ws.append([None if pd.isna(v) or v == "" else (round(v, 1) if isinstance(v, float) else v) for v in row])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions
    letters = {col: ws.cell(1, i + 1).column_letter for i, (col, _) in enumerate(SHEET_COLUMNS)}
    for col, values in (("unit_2026_27", (*UNITS, INJURED)), ("role1_2026_27", ROLES), ("role2_2026_27", ROLES)):
        dv = DataValidation(type="list", formula1=f'"{",".join(values)}"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{letters[col]}2:{letters[col]}{ws.max_row}")
    for col in INPUT_COLUMNS:
        ws.column_dimensions[letters[col]].width = 40 if col == "notes" else 15
        for r in range(1, ws.max_row + 1):
            ws[f"{letters[col]}{r}"].fill = INPUT_FILL
    ws.column_dimensions[letters["player"]].width = 24
    ws.column_dimensions[letters["yahoo_player_id"]].hidden = True
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_name(f".{path.name}.tmp")
    wb.save(tmp)
    tmp.replace(path)


def main(argv=None) -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--cached", action="store_true", help="reuse the cached tracking file")
    args = parser.parse_args(argv)
    if args.cached:
        tracking = pd.read_csv(dataset_path("tracking", SEASON, REGULAR))
        missing = [c for c in TRACKING_EXTRA if c not in tracking]
        if missing:
            sys.exit(f"ABORT: cached tracking predates {missing}; run without --cached")
    else:
        tracking = pull_tracking()
        if len(tracking) < MIN_TRACKING_ROWS:
            sys.exit(f"ABORT: tracking returned {len(tracking)} rows (< {MIN_TRACKING_ROWS}); sheet left untouched")
        write_dataset(tracking, "tracking", SEASON, REGULAR)
    lines = pd.read_csv(STATS_DIR / "season_lines" / f"{SEASON}_regular.csv")
    lines = lines[["nba_player_id", "TEAM_ABBREVIATION", "GP", "MIN", "PTS", "AST", "FGA", "FG3A", "BLK"]].merge(
        tracking, on="nba_player_id", how="left").merge(load_shot_totals(), on="nba_player_id", how="left")
    for total, per_game in (("FINISH_TOTAL", "FINISH_FGA"), ("AST_FGM_TOTAL", "AST_FGM"),
                            ("UNAST_FGM_TOTAL", "UNAST_FGM"), ("POA_TOTAL", "POA_MIN")):
        lines[per_game] = lines[total].fillna(0) / lines["GP"]
    lines["GS"] = lines["GS"].fillna(0)
    previous = read_previous(SHEET_PATH)
    legacy = is_legacy(previous)
    table = build_table(read_map(), classify(derive_signals(lines)), load_adp())
    table = carry_over(table, migrate_legacy(previous) if legacy else previous)
    if legacy:
        table = seed_roles(table)
        print("migrated the depth/role/option sheet: depth -> unit, untouched roles re-seeded, option dropped")
    write_sheet(table, SHEET_PATH)
    classified = table["role1"].notna().sum()
    print(f"{len(table)} players, {classified} with a 2025-26 role, "
          f"{(table['changed_team'] == 'Y').sum()} on a new team -> {SHEET_PATH}")


if __name__ == "__main__":
    main()
